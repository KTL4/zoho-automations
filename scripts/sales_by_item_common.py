"""
Shared logic for Zoho Inventory "Sales by Item" report automations.

Used by both item_runway_report.py (monthly, 14-week comparison, run on the
1st of each month) and item_runway_weekly_report.py (weekly, 36-week
comparison, run every Sunday). Both replicate Zoho Inventory's Reports >
Sales by Item; the only differences between them are the comparison window
length, the reference date used to anchor that window, and the output
filename pattern -- everything else (auth, rate limiting, catalog fetch,
invoice fetch, aggregation, filtering, Excel formatting) lives here so a
fix in one place (e.g. the rate limiter) benefits both automations instead
of needing to be duplicated and kept in sync by hand.

Zoho Inventory's API has no direct "Sales by Item" report endpoint (a Zoho
community thread confirms this gap: "this functionality exists in the
Reports interface, but there is no API available for it" -- same situation
as the SOH report's "Stock Summary"), so this is built from the underlying
transactional data instead:

  1. The full active item catalog is fetched (list + per-item detail, same
     approach as soh_daily_report.py) to get SKU / Item Name / Brand for
     every item, restricted to stock-tracked ("inventory" type) items --
     this drops non-sellable catalog entries like freight/container line
     items (e.g. "20FT Container") that aren't part of physical stock.
  2. Every non-draft, non-void invoice dated within the comparison window
     is listed (cheap, paginated, sorted newest-first so pagination can
     stop as soon as it walks past the window) and then fetched in full
     (one call per invoice, same as the per-item calls in the SOH report)
     to get its line items.
  3. Each line item's quantity is summed into the Monday-Sunday week its
     invoice date falls in.
  4. Items with zero sales across the entire window are dropped from the
     output -- the report only lists items that actually sold at least
     once in the window.

Both the catalog fetch and the invoice detail fetch hit the API far harder
than a naive loop can survive -- soh_daily_report.py's production run once
crashed after being 429'd for ~30 minutes straight -- so both phases share
the same rate limiter, heartbeat logging, and retry-with-backoff pattern
that fixed that.

Assumptions worth verifying against the real Zoho UI report before trusting
either automation (per the project's testing convention):
  - "Sales by Item" == invoiced quantity (draft/void invoices excluded),
    not booked sales-order quantity.
  - Report covers all warehouses/locations combined, not a single store.
  - Weeks run Monday-Sunday; each week's column header is that week's ISO
    week-of-year number (e.g. "WK 36"), oldest week leftmost.
  - The catalog is scoped to active items (Status.Active), matching the
    SOH report's convention, further restricted to item_type == "inventory"
    to exclude non-sellable/non-stock catalog entries. An item missing the
    item_type field entirely is kept rather than dropped, to fail open.

Required environment variables:
    ZOHO_CLIENT_ID
    ZOHO_CLIENT_SECRET
    ZOHO_REFRESH_TOKEN

Optional environment variables (read by each script's main(), not by this
module directly):
    ZOHO_ORGANIZATION_ID   - required if the Zoho account has more than one
                              organization (auto-detected otherwise)
    REPORT_TZ               - IANA timezone used to determine "today" (and
                               therefore the report window), defaults to
                               "Africa/Nairobi"
    REPORT_AS_OF_DATE       - overrides "today" with a fixed YYYY-MM-DD date,
                               for backfilling/testing what a run would have
                               produced on a specific past date. Unset in
                               normal scheduled runs.
    RATE_LIMIT_PER_MINUTE   - max Zoho API requests/minute, defaults to 60
"""
import os
import sys
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ACCOUNTS_TOKEN_URL = "https://accounts.zoho.com/oauth/v2/token"
STATIC_COLUMNS = ["SKU", "Item Name", "Brand"]
STOCK_ITEM_TYPE = "inventory"
EXCLUDED_INVOICE_STATUSES = {"draft", "void"}
INVOICE_LIST_PAGE_SIZE = 200
TOKEN_REFRESH_INTERVAL_SECONDS = 45 * 60  # access tokens expire after 1 hour
DETAIL_FETCH_WORKERS = 6
PROGRESS_EVERY = 100
HEARTBEAT_SECONDS = 20
RETRY_ATTEMPTS = 6
DEFAULT_RATE_LIMIT_PER_MINUTE = 60
MAX_RETRY_AFTER_SECONDS = 90


def get_access_token():
    response = requests.post(
        ACCOUNTS_TOKEN_URL,
        data={
            "grant_type": "refresh_token",
            "client_id": os.environ["ZOHO_CLIENT_ID"],
            "client_secret": os.environ["ZOHO_CLIENT_SECRET"],
            "refresh_token": os.environ["ZOHO_REFRESH_TOKEN"],
        },
        timeout=30,
    )
    response.raise_for_status()
    data = response.json()
    if "access_token" not in data:
        raise RuntimeError(f"Zoho did not return an access token: {data}")
    return data["access_token"], data["api_domain"]


class TokenStore:
    """Thread-safe access token holder that transparently refreshes long-running jobs."""

    def __init__(self, access_token):
        self._lock = threading.Lock()
        self._token = access_token
        self._issued_at = time.monotonic()

    def get(self):
        with self._lock:
            if time.monotonic() - self._issued_at > TOKEN_REFRESH_INTERVAL_SECONDS:
                self._token, _ = get_access_token()
                self._issued_at = time.monotonic()
            return self._token


class RateLimiter:
    """Paces requests across all worker threads to a fixed rate per minute.

    See soh_daily_report.py -- a live run there got 429'd for the better
    part of 30 minutes and crashed before this was added. Both the item
    catalog fetch and the invoice detail fetch below share one instance of
    this so the two phases don't each independently blow the budget.
    """

    def __init__(self, max_per_minute):
        self._lock = threading.Lock()
        self._interval = 60.0 / max_per_minute
        self._next_slot = time.monotonic()

    def wait(self):
        with self._lock:
            now = time.monotonic()
            start = max(now, self._next_slot)
            self._next_slot = start + self._interval
        sleep_for = start - now
        if sleep_for > 0:
            time.sleep(sleep_for)


def resolve_today(report_tz):
    """Returns the reference date for the report: REPORT_AS_OF_DATE if set
    (for backfilling/testing what a run would have produced on a past
    date), otherwise the current date in report_tz."""
    override = os.environ.get("REPORT_AS_OF_DATE")
    if override:
        return datetime.strptime(override, "%Y-%m-%d").date()
    return datetime.now(report_tz).date()


def resolve_organization_id(session, api_domain):
    org_id = os.environ.get("ZOHO_ORGANIZATION_ID")
    if org_id:
        return org_id

    response = session.get(f"{api_domain}/inventory/v1/organizations", timeout=30)
    response.raise_for_status()
    organizations = response.json().get("organizations", [])
    if len(organizations) == 1:
        return organizations[0]["organization_id"]

    names = ", ".join(f"{o['name']} ({o['organization_id']})" for o in organizations)
    raise RuntimeError(
        "Zoho account has multiple organizations; set ZOHO_ORGANIZATION_ID to "
        f"one of: {names}"
    )


def parse_retry_after(response, attempt):
    header_value = response.headers.get("Retry-After")
    if header_value is not None:
        try:
            return min(float(header_value), MAX_RETRY_AFTER_SECONDS)
        except ValueError:
            pass
    return min(2**attempt, MAX_RETRY_AFTER_SECONDS)


def fetch_detail(session, api_domain, organization_id, token_store, rate_limiter, url, result_key, label):
    """Shared retry/backoff/rate-limit logic for a single detail GET (item or invoice)."""
    last_error = None
    for attempt in range(RETRY_ATTEMPTS):
        rate_limiter.wait()
        try:
            response = session.get(
                url,
                params={"organization_id": organization_id},
                headers={"Authorization": f"Zoho-oauthtoken {token_store.get()}"},
                timeout=30,
            )
        except requests.exceptions.RequestException as exc:
            last_error = exc
            print(f"  [{label}] attempt {attempt + 1} network error: {exc}", file=sys.stderr, flush=True)
            time.sleep(min(2**attempt, MAX_RETRY_AFTER_SECONDS))
            continue

        if response.status_code == 429:
            wait_s = parse_retry_after(response, attempt)
            last_error = requests.exceptions.HTTPError(f"429 Too Many Requests for {label}")
            print(f"  [{label}] attempt {attempt + 1} rate-limited (429), waiting {wait_s:.1f}s", file=sys.stderr, flush=True)
            time.sleep(wait_s)
            continue

        try:
            response.raise_for_status()
        except requests.exceptions.HTTPError as exc:
            last_error = exc
            print(f"  [{label}] attempt {attempt + 1} HTTP error: {exc}", file=sys.stderr, flush=True)
            time.sleep(min(2**attempt, MAX_RETRY_AFTER_SECONDS))
            continue

        return response.json()[result_key]

    print(f"  [{label}] giving up after {RETRY_ATTEMPTS} attempts: {last_error}", file=sys.stderr, flush=True)
    return None


def fetch_concurrently(entity_ids, fetch_one, description):
    """Runs fetch_one(entity_id) across a thread pool with progress + heartbeat logging.

    Shared by both the item-catalog fetch and the invoice-detail fetch --
    see the module docstring for why this can't just be a naive loop.
    """
    results = []
    completed = 0
    failed = 0
    stop_heartbeat = threading.Event()

    def heartbeat():
        while not stop_heartbeat.wait(HEARTBEAT_SECONDS):
            print(
                f"  ...still working on {description}: {completed}/{len(entity_ids)} fetched ({failed} failed so far)",
                flush=True,
            )

    heartbeat_thread = threading.Thread(target=heartbeat, daemon=True)
    heartbeat_thread.start()

    try:
        with ThreadPoolExecutor(max_workers=DETAIL_FETCH_WORKERS) as executor:
            futures = {executor.submit(fetch_one, entity_id): entity_id for entity_id in entity_ids}
            for future in as_completed(futures):
                detail = future.result()
                completed += 1
                if detail is None:
                    failed += 1
                else:
                    results.append(detail)
                if completed % PROGRESS_EVERY == 0 or completed == len(entity_ids):
                    print(f"  ...{description}: {completed}/{len(entity_ids)} fetched ({failed} failed)", flush=True)
    finally:
        stop_heartbeat.set()
        heartbeat_thread.join()

    if failed:
        print(f"Warning: {failed} {description} fetch(es) failed after retries and were skipped.", file=sys.stderr, flush=True)

    return results


def fetch_active_item_ids(session, api_domain, organization_id):
    item_ids = []
    page = 1
    while True:
        response = session.get(
            f"{api_domain}/inventory/v1/items",
            params={
                "organization_id": organization_id,
                "filter_by": "Status.Active",
                "page": page,
                "per_page": 200,
            },
            timeout=30,
        )
        response.raise_for_status()
        payload = response.json()
        item_ids.extend(item["item_id"] for item in payload.get("items", []))

        if not payload.get("page_context", {}).get("has_more_page"):
            break
        page += 1

    return item_ids


def find_custom_field(item, *label_fragments):
    for field in item.get("custom_fields") or []:
        label = (field.get("label") or "").strip().lower()
        if any(fragment in label for fragment in label_fragments):
            value = field.get("value")
            if value:
                return value
    return ""


def extract_brand(item):
    return item.get("brand") or find_custom_field(item, "brand")


def is_stock_item(item):
    """True for physical, stock-tracked items -- excludes non-sellable
    catalog entries like freight/container line items (item_type values
    such as "purchases" or "service") that aren't part of physical stock
    and shouldn't appear in a stock-runway report. An item missing the
    item_type field entirely is kept rather than dropped, to fail open."""
    item_type = (item.get("item_type") or "").strip().lower()
    if not item_type:
        return True
    return item_type == STOCK_ITEM_TYPE


def build_item_catalog(session, api_domain, organization_id, token_store, rate_limiter):
    item_ids = fetch_active_item_ids(session, api_domain, organization_id)
    print(f"Found {len(item_ids)} active item(s); fetching item detail for SKU/Name/Brand...", flush=True)

    def fetch_one(item_id):
        return fetch_detail(
            session, api_domain, organization_id, token_store, rate_limiter,
            url=f"{api_domain}/inventory/v1/items/{item_id}",
            result_key="item",
            label=f"item {item_id}",
        )

    items = fetch_concurrently(item_ids, fetch_one, "item detail")

    catalog = {}
    skipped_non_stock = 0
    for item in items:
        if not is_stock_item(item):
            skipped_non_stock += 1
            continue
        catalog[item["item_id"]] = {
            "SKU": item.get("sku", ""),
            "Item Name": item.get("name", ""),
            "Brand": extract_brand(item),
        }

    if skipped_non_stock:
        print(
            f"Excluded {skipped_non_stock} non-stock item(s) (e.g. freight/container line items) "
            "from the catalog.",
            flush=True,
        )

    return catalog


def compute_week_buckets_ending_by(reference_date, num_periods):
    """Returns `num_periods` (start, end) date tuples, oldest first, for the
    `num_periods` Monday-Sunday weeks most recently completed on or before
    `reference_date` (inclusive -- if `reference_date` is itself a Sunday,
    the week ending that day is the most recent bucket)."""
    days_since_last_sunday = (reference_date.weekday() + 1) % 7  # Mon=0..Sun=6 -> Sun=0
    most_recent_week_end = reference_date - timedelta(days=days_since_last_sunday)

    buckets = []
    for i in range(num_periods):
        week_end = most_recent_week_end - timedelta(weeks=i)
        week_start = week_end - timedelta(days=6)
        buckets.append((week_start, week_end))
    buckets.reverse()
    return buckets


def week_label(week_start):
    return f"WK {week_start.isocalendar()[1]}"


def fetch_invoice_ids_in_window(session, api_domain, organization_id, window_start, window_end):
    """Lists invoice IDs dated within [window_start, window_end], excluding
    draft/void. Paginates newest-first and stops as soon as it walks past
    window_start, so this stays cheap regardless of total invoice history."""
    ids = []
    page = 1
    while True:
        response = session.get(
            f"{api_domain}/inventory/v1/invoices",
            params={
                "organization_id": organization_id,
                "sort_column": "date",
                "sort_order": "D",
                "page": page,
                "per_page": INVOICE_LIST_PAGE_SIZE,
            },
            timeout=30,
        )
        response.raise_for_status()
        payload = response.json()
        invoices = payload.get("invoices", [])

        reached_window_start = False
        for invoice in invoices:
            try:
                invoice_date = datetime.strptime(invoice["date"], "%Y-%m-%d").date()
            except (KeyError, ValueError):
                continue
            if invoice_date > window_end:
                continue
            if invoice_date < window_start:
                reached_window_start = True
                break
            if (invoice.get("status") or "").strip().lower() in EXCLUDED_INVOICE_STATUSES:
                continue
            ids.append(invoice["invoice_id"])

        if reached_window_start or not payload.get("page_context", {}).get("has_more_page"):
            break
        page += 1

    return ids


def aggregate_sales(invoice_details, buckets):
    """item_id -> list[len(buckets)] of quantity sold, one slot per bucket."""
    sales = defaultdict(lambda: [0] * len(buckets))
    for invoice in invoice_details:
        if invoice is None:
            continue
        try:
            invoice_date = datetime.strptime(invoice["date"], "%Y-%m-%d").date()
        except (KeyError, ValueError):
            continue

        bucket_idx = next(
            (idx for idx, (start, end) in enumerate(buckets) if start <= invoice_date <= end),
            None,
        )
        if bucket_idx is None:
            continue

        for line in invoice.get("line_items") or []:
            item_id = line.get("item_id")
            if not item_id:
                continue
            sales[item_id][bucket_idx] += line.get("quantity") or 0

    return sales


def build_rows(catalog, sales, buckets):
    rows = []
    for item_id, info in catalog.items():
        quantities = sales.get(item_id, [0] * len(buckets))
        if not any(quantities):
            continue  # no sales anywhere in the window -- omit from the report
        rows.append([info["SKU"], info["Item Name"], info["Brand"], *quantities])

    skipped_item_ids = set(sales) - set(catalog)
    if skipped_item_ids:
        print(
            f"Warning: {len(skipped_item_ids)} item(s) had sales in the window but are not in the "
            "active/stock-tracked item catalog (likely discontinued or non-stock); their sales were "
            "excluded from the report.",
            file=sys.stderr,
        )

    rows.sort(key=lambda row: row[1].lower())
    return rows


def write_excel(rows, week_columns, output_path):
    columns = STATIC_COLUMNS + week_columns

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(columns)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "B2"

    for row in rows:
        sheet.append(row)

    first_week_col = len(STATIC_COLUMNS) + 1
    for row_idx in range(2, len(rows) + 2):
        for col_idx in range(first_week_col, len(columns) + 1):
            sheet.cell(row=row_idx, column=col_idx).number_format = "#,##0"

    for col_idx, column_name in enumerate(columns, start=1):
        if col_idx < first_week_col:
            longest = max(
                [len(column_name)] + [len(str(row[col_idx - 1])) for row in rows],
                default=len(column_name),
            )
            width = longest + 4
        else:
            width = max(len(column_name) + 2, 8)
        sheet.column_dimensions[get_column_letter(col_idx)].width = width

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook.save(output_path)


def run_report(rate_limit_per_minute, buckets, output_path):
    """Shared end-to-end flow: auth, catalog fetch, invoice fetch/aggregate,
    write. `buckets` is whatever compute_week_buckets_ending_by returned."""
    window_start, window_end = buckets[0][0], buckets[-1][1]
    week_columns = [week_label(start) for start, _ in buckets]

    print(
        f"Report window: {window_start} to {window_end} "
        f"({len(buckets)} weekly periods, {week_columns[0]}..{week_columns[-1]})",
        flush=True,
    )

    rate_limiter = RateLimiter(rate_limit_per_minute)

    access_token, api_domain = get_access_token()
    token_store = TokenStore(access_token)
    session = requests.Session()
    session.headers["Authorization"] = f"Zoho-oauthtoken {access_token}"

    organization_id = resolve_organization_id(session, api_domain)

    catalog = build_item_catalog(session, api_domain, organization_id, token_store, rate_limiter)

    invoice_ids = fetch_invoice_ids_in_window(session, api_domain, organization_id, window_start, window_end)
    print(f"Found {len(invoice_ids)} invoice(s) in window; fetching line item detail...", flush=True)

    def fetch_one_invoice(invoice_id):
        return fetch_detail(
            session, api_domain, organization_id, token_store, rate_limiter,
            url=f"{api_domain}/inventory/v1/invoices/{invoice_id}",
            result_key="invoice",
            label=f"invoice {invoice_id}",
        )

    invoice_details = fetch_concurrently(invoice_ids, fetch_one_invoice, "invoice detail")

    sales = aggregate_sales(invoice_details, buckets)
    rows = build_rows(catalog, sales, buckets)

    write_excel(rows, week_columns, output_path)
    print(f"Wrote {len(rows)} rows ({len(week_columns)} weekly columns) to {output_path}")
