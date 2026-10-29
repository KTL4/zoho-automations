"""
Container Sellout Report -- how much of every container that has ever
arrived has sold, and how fast.

Zoho Inventory doesn't tag a sale with the container its stock came from,
so sales have to be matched to containers by FIFO (oldest stock sells
first). Zoho already does that matching itself in Reports > FIFO Cost Lot
Tracking: every inward transaction (a container bill, opening stock, an
adjustment, a transfer in) is a "lot" per SKU, with the quantity that came
in and the quantity of that lot still on hand. This report is built on
that, in three phases:

  1. Lots -- the whole FIFO Cost Lot Tracking report is pulled
     (`reports/fifocostlottracking`, note the "lot"; `fifocosttracking`
     404s). A lot is a *container lot* when it came in on a bill / purchase
     receive from one of the container suppliers (CONTAINER_SUPPLIER_PATTERNS,
     matched against the lot's contact name). Lots are grouped into one
     container per bill number. Per container this already gives the
     headline figure, straight from Zoho's own FIFO:
         sellout % = (units in - units still on hand) / units in
     Note that "units in - on hand" counts *every* way stock left, sales
     but also transfers out (e.g. KTL -> KSL) and write-offs.

  2. Sales -- every non-draft, non-void invoice (and, where the API exposes
     them, credit notes, netted off as returns) since the first lot, at
     line level (date, SKU, qty). This is the expensive phase -- one API
     call per document -- so the lines are cached in data/ and committed:
     each run lists the documents (cheap, 200 per call) and only fetches
     detail for ones that are new or whose last_modified_time changed.
     A run that hits SALES_FETCH_TIME_BUDGET_MINUTES saves what it has and
     the next run carries on, so the very first backfill can span several
     runs without ever losing work.

  3. Timing -- the net daily sales of each SKU are matched to its lots
     oldest-first. Each lot can only take as many sales as Zoho says left
     it (units in - on hand), so the timing view never contradicts Zoho's
     own FIFO. What Zoho consumed from a lot that sales don't account for
     is reported as "other outflows" (transfers, adjustments, damages).
     This gives, per container: % sold within 30/60/90/180 days of arrival
     and the days it took to sell 50% and 80% of it.

Assumptions worth verifying against the Zoho UI before trusting a run (per
the project's testing convention):
  - The FIFO Cost Lot Tracking API returns one row per inward lot with
    date (YYYY-MM-DD), transaction_type, transaction_number, contact_name,
    sku, quantity_in, quantity_on_hand, cost_price_per_unit -- the same
    fields the inventory-workflows toolkit's pull_fifo_lots.ps1 reads.
  - A container is one bill; its bill number is the container's name.
  - Sales are matched to lots by SKU (case-insensitive), not item_id, since
    the lot report is keyed on SKU.
  - The timing view is FIFO by date only: within a SKU, a sale is matched
    to the oldest lot that had arrived by the sale date and still had
    room. Transfers/adjustments are not dated in this model, so they are
    assumed not to change which lot a sale came from.

Required environment variables:
    ZOHO_CLIENT_ID, ZOHO_CLIENT_SECRET, ZOHO_REFRESH_TOKEN

Optional environment variables:
    ZOHO_ORGANIZATION_ID              - required if the account has more than
                                         one organization
    REPORT_TZ                          - defaults to "Africa/Nairobi"
    REPORT_AS_OF_DATE                  - overrides "today" (YYYY-MM-DD)
    LOTS_FROM_DATE                     - earliest lot date to pull, defaults
                                         to 2015-01-01 (i.e. all history)
    CONTAINER_SUPPLIER_PATTERNS        - comma-separated, case-insensitive
                                         fragments of a lot's contact name
                                         that mark it as a container supplier
    SALES_FETCH_TIME_BUDGET_MINUTES    - stop fetching sales detail after this
                                         long and save progress, default 300
    RATE_LIMIT_PER_MINUTE              - defaults to 60
"""
import csv
import os
import sys
import time
from collections import defaultdict
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sales_by_item_common as common

DEFAULT_LOTS_FROM_DATE = "2015-01-01"
DEFAULT_SUPPLIER_PATTERNS = "UNCLE BILL,UNCLEBILL,UBL,AIWIBI,AIWINA,YIFAR,EFRIS STOCK"
CONTAINER_TRANSACTION_TYPES = ("bill", "purchase receive", "purchasereceive", "receive")
SOLD_BY_DAYS = (30, 60, 90, 180)
SOLD_SHARE_MILESTONES = (0.5, 0.8)
LOT_PAGE_SIZE = 200
LOT_MAX_PAGES = 500
LOT_RETRY_WAITS = (5, 8, 12, 20, 30, 45, 60, 60)
SALES_CHUNK_SIZE = 300
DEFAULT_TIME_BUDGET_MINUTES = 300
DATA_DIR = "data"
CACHE_COLUMNS = ["doc_type", "doc_id", "doc_date", "last_modified_time", "sku", "item_id", "quantity"]

# (doc_type, list path, list key, id field, detail key). Invoices count as
# sales, credit notes are subtracted as returns.
SALES_DOCUMENTS = [
    ("invoice", "invoices", "invoices", "invoice_id", "invoice"),
    ("creditnote", "creditnotes", "creditnotes", "creditnote_id", "creditnote"),
]


def to_number(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    for prefix in ("UGX", "KES", "USD", "Ugx", "Ksh"):
        if text.startswith(prefix):
            text = text[len(prefix):].strip()
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_date(value):
    text = str(value or "").strip()
    for candidate, fmt in ((text[:10], "%Y-%m-%d"), (text, "%d %b %Y"), (text[:10], "%d/%m/%Y")):
        try:
            return datetime.strptime(candidate, fmt).date()
        except ValueError:
            continue
    return None


def sku_key(sku):
    return str(sku or "").strip().upper()


# ---------------------------------------------------------------- phase 1: lots

def fetch_lot_page(session, api_domain, organization_id, from_date, to_date, page):
    """One page of the FIFO Cost Lot Tracking report, waiting out Zoho's
    "still computing" state and 429s. The inventory-workflows toolkit found
    that this report can serve a stale ThisMonth default instead of the
    requested range when the access token is reused, so a suspiciously
    short single page triggers a fresh token and a retry."""
    params = {
        "organization_id": organization_id,
        "filter_by": "TransactionDate.CustomDate",
        "from_date": from_date.isoformat(),
        "to_date": to_date.isoformat(),
        "page": page,
        "per_page": LOT_PAGE_SIZE,
    }
    url = f"{api_domain}/inventory/v1/reports/fifocostlottracking"
    for attempt, wait_s in enumerate(LOT_RETRY_WAITS):
        response = session.get(url, params=params, timeout=120)
        if response.status_code == 429:
            print(f"  lots page {page}: 429, waiting {wait_s}s", flush=True)
            time.sleep(wait_s)
            continue
        response.raise_for_status()
        payload = response.json()
        rows = payload.get("fifo_cost_lot_tracking") or []
        context = payload.get("page_context") or {}
        in_progress = context.get("is_inv_txn_job_in_progress") is True
        looks_stale = page == 1 and not context.get("has_more_page") and len(rows) < 50
        if (in_progress or looks_stale) and attempt < len(LOT_RETRY_WAITS) - 1:
            print(f"  lots page {page}: range not ready ({len(rows)} rows), fresh token + retry in {wait_s}s", flush=True)
            time.sleep(wait_s)
            access_token, _ = common.get_access_token()
            session.headers["Authorization"] = f"Zoho-oauthtoken {access_token}"
            continue
        return rows, bool(context.get("has_more_page"))
    raise RuntimeError(f"Gave up on FIFO lot report page {page}")


def fetch_lots(session, api_domain, organization_id, from_date, to_date):
    lots = []
    page = 1
    more = True
    while more and page <= LOT_MAX_PAGES:
        rows, more = fetch_lot_page(session, api_domain, organization_id, from_date, to_date, page)
        for row in rows:
            lot = normalize_lot(row, len(lots))
            if lot is not None:
                lots.append(lot)
        if page % 10 == 0:
            print(f"  ...lots: {page} page(s), {len(lots)} inward lot(s) so far", flush=True)
        page += 1
        time.sleep(1.0)
    if more:
        print(f"Warning: stopped after {LOT_MAX_PAGES} pages of lots; history is truncated.", file=sys.stderr, flush=True)
    return lots


def normalize_lot(row, sequence):
    """A raw report row -> lot dict, or None for rows that aren't inward lots."""
    quantity_in = to_number(row.get("quantity_in"))
    lot_date = parse_date(row.get("date"))
    if quantity_in <= 0 or lot_date is None:
        return None
    on_hand = min(max(to_number(row.get("quantity_on_hand")), 0.0), quantity_in)
    unit_cost = to_number(row.get("cost_price_per_unit"))
    return {
        "seq": sequence,
        "date": lot_date,
        "transaction_type": str(row.get("transaction_type") or "").strip(),
        "transaction_number": str(row.get("transaction_number") or "").strip(),
        "contact_name": str(row.get("contact_name") or "").strip(),
        "lot_number": str(row.get("lot_number") or "").strip(),
        "sku": str(row.get("sku") or "").strip(),
        "item_name": str(row.get("item_name") or "").strip(),
        "quantity_in": quantity_in,
        "quantity_on_hand": on_hand,
        "unit_cost": unit_cost,
    }


def supplier_patterns():
    raw = os.environ.get("CONTAINER_SUPPLIER_PATTERNS") or DEFAULT_SUPPLIER_PATTERNS
    return [p.strip().upper() for p in raw.split(",") if p.strip()]


def supplier_of(contact_name, patterns):
    """Groups a lot's contact into a supplier family, or None when it isn't
    a container supplier."""
    name = contact_name.upper()
    if not any(p in name for p in patterns):
        return None
    if "UNCLE" in name or name.startswith("UBL"):
        return "UBL"
    if "AIWIBI" in name or "AIWINA" in name or "YIFAR" in name:
        return "Aiwibi"
    return contact_name


def is_container_lot(lot, patterns):
    ttype = lot["transaction_type"].lower()
    return any(t in ttype for t in CONTAINER_TRANSACTION_TYPES) and supplier_of(lot["contact_name"], patterns) is not None


# ---------------------------------------------------------------- phase 2: sales

def list_documents(session, api_domain, organization_id, path, list_key, id_field, since):
    """{doc_id: (date, last_modified_time)} for non-draft, non-void documents
    dated on/after `since`. Newest-first so pagination stops at `since`."""
    docs = {}
    page = 1
    while True:
        response = session.get(
            f"{api_domain}/inventory/v1/{path}",
            params={
                "organization_id": organization_id,
                "sort_column": "date",
                "sort_order": "D",
                "page": page,
                "per_page": common.INVOICE_LIST_PAGE_SIZE,
            },
            timeout=60,
        )
        if response.status_code == 429:
            time.sleep(common.parse_retry_after(response, 3))
            continue
        response.raise_for_status()
        payload = response.json()
        reached_start = False
        for doc in payload.get(list_key) or []:
            doc_date = parse_date(doc.get("date"))
            if doc_date is None:
                continue
            if doc_date < since:
                reached_start = True
                break
            if (doc.get("status") or "").strip().lower() in common.EXCLUDED_INVOICE_STATUSES:
                continue
            if doc.get(id_field):
                docs[str(doc[id_field])] = (doc_date.isoformat(), str(doc.get("last_modified_time") or ""))
        if reached_start or not (payload.get("page_context") or {}).get("has_more_page"):
            break
        page += 1
    return docs


def cache_path(organization_id):
    return os.path.join(DATA_DIR, f"sales_lines_{organization_id}.csv")


def load_cache(path):
    """{(doc_type, doc_id): {"date", "modified", "lines": [(sku, item_id, qty)]}}"""
    cache = {}
    if not os.path.exists(path):
        return cache
    with open(path, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            key = (row["doc_type"], row["doc_id"])
            entry = cache.setdefault(key, {"date": row["doc_date"], "modified": row["last_modified_time"], "lines": []})
            if row["sku"] or row["item_id"]:
                entry["lines"].append((row["sku"], row["item_id"], to_number(row["quantity"])))
    return cache


def save_cache(path, cache):
    """Sorted by date so a weekly refresh mostly appends to the file."""
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(CACHE_COLUMNS)
        for (doc_type, doc_id), entry in sorted(cache.items(), key=lambda kv: (kv[1]["date"], kv[0])):
            lines = entry["lines"] or [("", "", 0)]  # placeholder so an empty document stays "fetched"
            for sku, item_id, qty in lines:
                writer.writerow([doc_type, doc_id, entry["date"], entry["modified"], sku, item_id, qty])
    os.replace(tmp, path)


def document_lines(detail):
    lines = []
    for line in detail.get("line_items") or []:
        qty = to_number(line.get("quantity"))
        if qty and (line.get("sku") or line.get("item_id")):
            lines.append((str(line.get("sku") or ""), str(line.get("item_id") or ""), qty))
    return lines


def sync_sales(session, api_domain, organization_id, token_store, rate_limiter, since, deadline):
    """Brings the cached sales lines up to date. Returns (cache, complete)."""
    path = cache_path(organization_id)
    cache = load_cache(path)
    complete = True

    for doc_type, list_path, list_key, id_field, detail_key in SALES_DOCUMENTS:
        # The detail fetches below can run for hours; listings use the
        # session's header, so give it a current token each time.
        session.headers["Authorization"] = f"Zoho-oauthtoken {token_store.get()}"
        try:
            listed = list_documents(session, api_domain, organization_id, list_path, list_key, id_field, since)
        except requests.exceptions.HTTPError as exc:
            if doc_type == "invoice":
                raise
            print(f"Warning: could not list {list_path} ({exc}); returns are not netted off.", file=sys.stderr, flush=True)
            continue

        # Drop cached documents that were since voided/deleted or edited.
        for key in [k for k in cache if k[0] == doc_type]:
            current = listed.get(key[1])
            if current is None or current[1] != cache[key]["modified"]:
                del cache[key]

        todo = sorted((doc_id for doc_id in listed if (doc_type, doc_id) not in cache), key=lambda d: listed[d][0])
        print(f"{list_path}: {len(listed)} since {since}, {len(todo)} to fetch ({len(listed) - len(todo)} cached)", flush=True)

        def fetch_one(doc_id, list_path=list_path, detail_key=detail_key):
            detail = common.fetch_detail(
                session, api_domain, organization_id, token_store, rate_limiter,
                url=f"{api_domain}/inventory/v1/{list_path}/{doc_id}",
                result_key=detail_key,
                label=f"{detail_key} {doc_id}",
            )
            return None if detail is None else (doc_id, detail)

        for start in range(0, len(todo), SALES_CHUNK_SIZE):
            if time.monotonic() > deadline:
                print(f"Time budget reached with {len(todo) - start} {list_path} still to fetch; "
                      "saving progress, the next run carries on.", flush=True)
                complete = False
                break
            chunk = todo[start:start + SALES_CHUNK_SIZE]
            for doc_id, detail in common.fetch_concurrently(chunk, fetch_one, f"{detail_key} detail"):
                doc_date, modified = listed[doc_id]
                cache[(doc_type, doc_id)] = {"date": doc_date, "modified": modified, "lines": document_lines(detail)}
            save_cache(path, cache)
            if any((doc_type, d) not in cache for d in chunk):
                complete = False  # some fetches failed after retries; retried next run

        save_cache(path, cache)
        if not complete:
            break

    return cache, complete


def daily_net_sales(cache):
    """{SKU: {date: net qty}} -- invoices minus credit notes. Lines without
    a SKU (e.g. freight or service lines) can't be matched to a lot."""
    sales = defaultdict(lambda: defaultdict(float))
    for (doc_type, _), entry in cache.items():
        sign = -1.0 if doc_type == "creditnote" else 1.0
        doc_date = parse_date(entry["date"])
        for sku, _, qty in entry["lines"]:
            if sku_key(sku):
                sales[sku_key(sku)][doc_date] += sign * qty
    return sales


# ---------------------------------------------------------------- phase 3: timing

def allocate_sales_fifo(lots, daily_sales):
    """Matches one SKU's net daily sales to its lots, oldest first.

    `lots` are this SKU's lot dicts; each can take at most what Zoho says
    left it (quantity_in - quantity_on_hand). A sale only goes to a lot that
    had arrived by the sale date. Returns ({lot seq: [(date, qty), ...]},
    unmatched qty). Net returns (negative days) are carried forward and
    netted off the next sales."""
    ordered = sorted(lots, key=lambda l: (l["date"], l["seq"]))
    room = [l["quantity_in"] - l["quantity_on_hand"] for l in ordered]
    allocations = defaultdict(list)
    unmatched = 0.0
    carry = 0.0
    for sale_date in sorted(daily_sales):
        qty = daily_sales[sale_date] + carry
        carry = 0.0
        if qty <= 0:
            carry = qty
            continue
        for idx, lot in enumerate(ordered):
            if qty <= 0:
                break
            if lot["date"] > sale_date:
                break
            take = min(room[idx], qty)
            if take > 0:
                room[idx] -= take
                qty -= take
                allocations[lot["seq"]].append((sale_date, take))
        unmatched += max(qty, 0.0)
    return allocations, unmatched


def summarize(lots, as_of):
    """Shared container / container x SKU metrics for a group of lots that
    already carry their "sales" allocation list."""
    arrival = min(l["date"] for l in lots)
    units_in = sum(l["quantity_in"] for l in lots)
    on_hand = sum(l["quantity_on_hand"] for l in lots)
    consumed = units_in - on_hand
    sales = sorted(((d - arrival).days, q) for l in lots for d, q in l["sales"])
    sold = sum(q for _, q in sales)
    age = (as_of - arrival).days

    sold_by = {}
    for days in SOLD_BY_DAYS:
        sold_by[days] = (sum(q for d, q in sales if d <= days) / units_in) if age >= days and units_in else None

    days_to = {}
    for share in SOLD_SHARE_MILESTONES:
        running, days_to[share] = 0.0, None
        for d, q in sales:
            running += q
            if running >= share * units_in - 1e-9:
                days_to[share] = d
                break

    return {
        "arrival": arrival,
        "age": age,
        "units_in": units_in,
        "on_hand": on_hand,
        "consumed": consumed,
        "sellout": consumed / units_in if units_in else 0.0,
        "sold": sold,
        "sold_pct": sold / units_in if units_in else 0.0,
        "other_out": max(consumed - sold, 0.0),
        "sold_by": sold_by,
        "days_to": days_to,
        "last_sale": (arrival + timedelta(days=sales[-1][0])) if sales else None,
        "cost_in": sum(l["quantity_in"] * l["unit_cost"] for l in lots),
        "cost_on_hand": sum(l["quantity_on_hand"] * l["unit_cost"] for l in lots),
    }


def build_report_data(lots, daily_sales, patterns, as_of):
    lots_by_sku = defaultdict(list)
    for lot in lots:
        lot["sales"] = []
        lots_by_sku[sku_key(lot["sku"])].append(lot)

    unmatched_total = 0.0
    for key, sku_lots in lots_by_sku.items():
        allocations, unmatched = allocate_sales_fifo(sku_lots, daily_sales.get(key, {}))
        unmatched_total += unmatched
        for lot in sku_lots:
            lot["sales"] = allocations.get(lot["seq"], [])

    containers = defaultdict(list)
    for lot in lots:
        if is_container_lot(lot, patterns):
            containers[(supplier_of(lot["contact_name"], patterns), lot["transaction_number"])].append(lot)

    container_rows, sku_rows = [], []
    for (supplier, number), c_lots in containers.items():
        container_rows.append({"supplier": supplier, "container": number, "skus": len({sku_key(l["sku"]) for l in c_lots}),
                               **summarize(c_lots, as_of)})
        by_sku = defaultdict(list)
        for lot in c_lots:
            by_sku[sku_key(lot["sku"])].append(lot)
        for s_lots in by_sku.values():
            sku_rows.append({"supplier": supplier, "container": number, "sku": s_lots[0]["sku"],
                             "item_name": s_lots[0]["item_name"], **summarize(s_lots, as_of)})

    container_rows.sort(key=lambda r: (r["arrival"], r["container"]))
    sku_rows.sort(key=lambda r: (r["arrival"], r["container"], r["sku"].lower()))
    return container_rows, sku_rows, unmatched_total


# ---------------------------------------------------------------- excel

def metric_cells(r):
    return [
        r["arrival"], r["age"], r["units_in"], r["sold"], r["other_out"], r["on_hand"],
        r["sellout"], r["sold_pct"],
        *[r["sold_by"][d] for d in SOLD_BY_DAYS],
        *[r["days_to"][s] for s in SOLD_SHARE_MILESTONES],
        r["last_sale"], r["cost_in"], r["cost_on_hand"],
    ]


METRIC_HEADERS = (
    ["Arrived", "Days Since Arrival", "Units In", "Units Sold", "Other Outflows", "Units On Hand",
     "Sellout % (Zoho FIFO)", "Sold % (Invoices)"]
    + [f"Sold % by Day {d}" for d in SOLD_BY_DAYS]
    + [f"Days to {int(s * 100)}% Sold" for s in SOLD_SHARE_MILESTONES]
    + ["Last Sale", "Cost In", "Cost On Hand"]
)
INT_HEADERS = {"Days Since Arrival", "Units In", "Units Sold", "Other Outflows", "Units On Hand", "SKUs"} | {
    f"Days to {int(s * 100)}% Sold" for s in SOLD_SHARE_MILESTONES}
MONEY_HEADERS = {"Cost In", "Cost On Hand"}
DATE_HEADERS = {"Arrived", "Last Sale"}


def write_sheet(sheet, headers, rows):
    sheet.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in rows:
        sheet.append(row)
    for col_idx, name in enumerate(headers, start=1):
        if name in DATE_HEADERS:
            fmt = "dd-mmm-yyyy"
        elif name in INT_HEADERS:
            fmt = "#,##0"
        elif name in MONEY_HEADERS:
            fmt = "#,##0"
        elif "%" in name:
            fmt = "0.0%"
        else:
            fmt = None
        if fmt:
            for row_idx in range(2, len(rows) + 2):
                sheet.cell(row=row_idx, column=col_idx).number_format = fmt
        longest = max([min(len(name), 14)] + [len(str(r[col_idx - 1])) for r in rows[:500]])
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(longest + 3, 10), 45)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = sheet.cell(row=2, column=3)
    sheet.auto_filter.ref = sheet.dimensions


def write_excel(container_rows, sku_rows, notes, output_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Containers"
    write_sheet(sheet, ["Container", "Supplier", "SKUs"] + METRIC_HEADERS,
                [[r["container"], r["supplier"], r["skus"], *metric_cells(r)] for r in container_rows])

    sheet = workbook.create_sheet("Container x SKU")
    write_sheet(sheet, ["Container", "Supplier", "SKU", "Item Name"] + METRIC_HEADERS,
                [[r["container"], r["supplier"], r["sku"], r["item_name"], *metric_cells(r)] for r in sku_rows])

    sheet = workbook.create_sheet("Read Me")
    sheet.column_dimensions["A"].width = 120
    for line in notes:
        sheet.append([line])
        sheet.cell(row=sheet.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A1"].font = Font(bold=True, size=13)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook.save(output_path)


def build_notes(as_of, organization_id, lots_from, sales_complete, unmatched, n_lots, n_container_lots):
    notes = [
        f"Container Sellout Report -- as of {as_of:%d %b %Y} (Zoho organization {organization_id})",
        "",
        "One row per container (one bill from a container supplier) that has ever arrived.",
        "Sellout % (Zoho FIFO) = (units in - units still on hand) / units in, straight from Zoho's FIFO Cost Lot "
        "Tracking report. It counts every way stock left: sales, but also transfers out and write-offs.",
        "Sold % (Invoices) = units invoiced (net of credit notes) that FIFO matches to this container / units in.",
        "Other Outflows = what Zoho shows as gone from the container that invoices don't account for "
        "(transfers to other orgs/stores, adjustments, damages).",
        "Sold % by Day N = share of the container invoiced within N days of arrival. Blank if the container "
        "is younger than N days.",
        "Days to 50% / 80% Sold = days after arrival until that share was invoiced. Blank if not reached yet.",
        "Cost In / Cost On Hand are in the organization's base currency at the lot's FIFO unit cost.",
        "",
        "How sales are matched to containers: Zoho doesn't record which container a sale came from, so each "
        "SKU's sales are matched oldest-container-first (FIFO), never more per container than Zoho's own FIFO "
        "says left it, and never before the container arrived.",
        "",
        f"Lots pulled from {lots_from} onward: {n_lots:,} inward lots, {n_container_lots:,} of them container lots.",
        f"Sales that FIFO couldn't place in any lot (no stock on record at the time): {unmatched:,.0f} units.",
    ]
    if not sales_complete:
        notes += ["", "PARTIAL: the sales history backfill hasn't finished yet, so Sold %, Sold % by Day N and "
                      "Days to N% are understated. Sellout % (Zoho FIFO) is complete. The next run carries on."]
    return notes


# ---------------------------------------------------------------- main

def save_lots(organization_id, lots):
    path = os.path.join(DATA_DIR, f"fifo_lots_{organization_id}.csv")
    os.makedirs(DATA_DIR, exist_ok=True)
    columns = ["date", "transaction_type", "transaction_number", "contact_name", "lot_number",
               "sku", "item_name", "quantity_in", "quantity_on_hand", "unit_cost"]
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(columns)
        for lot in sorted(lots, key=lambda l: (l["date"], l["transaction_number"], l["sku"])):
            writer.writerow([lot[c].isoformat() if c == "date" else lot[c] for c in columns])


def main():
    started = time.monotonic()
    report_tz = ZoneInfo(os.environ.get("REPORT_TZ") or "Africa/Nairobi")
    as_of = common.resolve_today(report_tz)
    lots_from = datetime.strptime(os.environ.get("LOTS_FROM_DATE") or DEFAULT_LOTS_FROM_DATE, "%Y-%m-%d").date()
    budget_min = float(os.environ.get("SALES_FETCH_TIME_BUDGET_MINUTES") or DEFAULT_TIME_BUDGET_MINUTES)
    rate_limiter = common.RateLimiter(int(os.environ.get("RATE_LIMIT_PER_MINUTE") or common.DEFAULT_RATE_LIMIT_PER_MINUTE))
    patterns = supplier_patterns()

    access_token, api_domain = common.get_access_token()
    token_store = common.TokenStore(access_token)
    session = requests.Session()
    session.headers["Authorization"] = f"Zoho-oauthtoken {access_token}"
    organization_id = common.resolve_organization_id(session, api_domain)

    print(f"Phase 1: FIFO lots {lots_from}..{as_of}", flush=True)
    lots = fetch_lots(session, api_domain, organization_id, lots_from, as_of)
    if not lots:
        raise RuntimeError("FIFO Cost Lot Tracking returned no inward lots; nothing to report.")
    save_lots(organization_id, lots)
    container_lot_count = sum(1 for l in lots if is_container_lot(l, patterns))
    first_lot = min(l["date"] for l in lots)
    print(f"{len(lots)} inward lot(s) from {first_lot}, {container_lot_count} container lot(s)", flush=True)
    types = defaultdict(int)
    for lot in lots:
        types[(lot["transaction_type"], supplier_of(lot["contact_name"], patterns) or "-")] += 1
    for (ttype, supplier), n in sorted(types.items(), key=lambda kv: -kv[1])[:15]:
        print(f"  {n:6d} lots  type={ttype!r:28} supplier={supplier}", flush=True)

    print(f"Phase 2: sales since {first_lot}", flush=True)
    deadline = started + budget_min * 60
    cache, sales_complete = sync_sales(session, api_domain, organization_id, token_store, rate_limiter, first_lot, deadline)
    daily_sales = daily_net_sales(cache)

    print("Phase 3: matching sales to containers", flush=True)
    container_rows, sku_rows, unmatched = build_report_data(lots, daily_sales, patterns, as_of)
    notes = build_notes(as_of, organization_id, lots_from, sales_complete, unmatched, len(lots), container_lot_count)

    output_path = os.path.join("reports", f"Container_Sellout_{as_of:%d_%m_%Y}.xlsx")
    write_excel(container_rows, sku_rows, notes, output_path)
    print(f"Wrote {len(container_rows)} container(s), {len(sku_rows)} container x SKU row(s) to {output_path}"
          f"{'' if sales_complete else ' (PARTIAL: sales backfill continues next run)'}", flush=True)


if __name__ == "__main__":
    main()
