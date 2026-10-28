"""
SOH Daily Report

Pulls current stock-on-hand for a given warehouse from Zoho Inventory and
writes it to an Excel file named SOH_DD_MM_YYYY.xlsx.

Zoho Inventory's API has no direct "Stock Summary" report endpoint, so this
approximates it using the Items API: the items *list* endpoint doesn't
include per-warehouse stock, so active items' full detail (which does
include a `warehouses` array) is fetched in batches of up to 100 via the
bulk /itemdetails endpoint, keeping the entry matching WAREHOUSE_NAME.
Batch fetches run concurrently, paced by a shared rate limiter to stay
under Zoho's request quota (a live run against ~1850 items crashed with a
429 after being throttled for the better part of 30 minutes back when this
fetched one item per call, which is what the rate limiter and Retry-After
handling below are for). If a whole batch fails after retries, its items
are re-fetched one at a time via the single-item endpoint instead, so one
bad id doesn't cost the rest of the batch. Progress prints after every
batch plus a heartbeat every 20 seconds (so a stalled run is visible in
real time rather than going silent), and any item that still fails after
retries is skipped rather than crashing the whole report. The access
token is refreshed automatically if a run takes long enough to approach
the 1-hour token expiry.

Required environment variables:
    ZOHO_CLIENT_ID
    ZOHO_CLIENT_SECRET
    ZOHO_REFRESH_TOKEN

Optional environment variables:
    ZOHO_ORGANIZATION_ID   - required if the Zoho account has more than one
                              organization (auto-detected otherwise)
    WAREHOUSE_NAME          - defaults to "Store 1"
    REPORT_TZ               - IANA timezone for naming the output file,
                               defaults to "Africa/Nairobi"
    RATE_LIMIT_PER_MINUTE   - max Zoho API requests/minute, defaults to 60
"""
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ACCOUNTS_TOKEN_URL = "https://accounts.zoho.com/oauth/v2/token"
OUTPUT_COLUMNS = ["BAR CODE", "SKU", "Item Name", "SOH", "Sales Price", "Brand"]
BARCODE_FIELDS = ("upc", "ean", "isbn", "part_number")
TOKEN_REFRESH_INTERVAL_SECONDS = 45 * 60  # access tokens expire after 1 hour
DETAIL_FETCH_WORKERS = 6
HEARTBEAT_SECONDS = 20
RETRY_ATTEMPTS = 6
DEFAULT_RATE_LIMIT_PER_MINUTE = 60
MAX_RETRY_AFTER_SECONDS = 90
ITEMDETAILS_BATCH_SIZE = 100  # Zoho's documented max ids per /itemdetails call


def get_access_token():
    response = requests.post(
        ACCOUNTS_TOKEN_URL,
        data={
            "grant_type": "refresh_token",
            "client_id": os.environ["ZOHO_CLIENT_ID"],
            "client_secret": os.environ["ZOHO_CLIENT_SECRET"],
            "refresh_token": os.environ["ZOHO_REFRESH_TOKEN"],
        },
        timeout=30,
    )
    response.raise_for_status()
    data = response.json()
    if "access_token" not in data:
        raise RuntimeError(f"Zoho did not return an access token: {data}")
    return data["access_token"], data["api_domain"]


class TokenStore:
    """Thread-safe access token holder that transparently refreshes long-running jobs."""

    def __init__(self, access_token):
        self._lock = threading.Lock()
        self._token = access_token
        self._issued_at = time.monotonic()

    def get(self):
        with self._lock:
            if time.monotonic() - self._issued_at > TOKEN_REFRESH_INTERVAL_SECONDS:
                self._token, _ = get_access_token()
                self._issued_at = time.monotonic()
            return self._token


class RateLimiter:
    """Paces requests across all worker threads to a fixed rate per minute.

    A short per-item exponential backoff isn't enough to recover from a
    sustained rate limit (a live run got 429'd for the better part of 30
    minutes and then crashed) -- this caps how fast requests actually go
    out in the first place.
    """

    def __init__(self, max_per_minute):
        self._lock = threading.Lock()
        self._interval = 60.0 / max_per_minute
        self._next_slot = time.monotonic()

    def wait(self):
        with self._lock:
            now = time.monotonic()
            start = max(now, self._next_slot)
            self._next_slot = start + self._interval
        sleep_for = start - now
        if sleep_for > 0:
            time.sleep(sleep_for)


def resolve_organization_id(session, api_domain):
    org_id = os.environ.get("ZOHO_ORGANIZATION_ID")
    if org_id:
        return org_id

    response = session.get(f"{api_domain}/inventory/v1/organizations", timeout=30)
    response.raise_for_status()
    organizations = response.json().get("organizations", [])
    if len(organizations) == 1:
        return organizations[0]["organization_id"]

    names = ", ".join(f"{o['name']} ({o['organization_id']})" for o in organizations)
    raise RuntimeError(
        "Zoho account has multiple organizations; set ZOHO_ORGANIZATION_ID to "
        f"one of: {names}"
    )


def fetch_active_item_ids(session, api_domain, organization_id):
    """Returns (item_ids, brand_by_id). The items *list* endpoint already
    includes each item's `brand` for free, which matters because the bulk
    /itemdetails endpoint used afterwards for per-warehouse stock does not
    return `brand` at all (confirmed 2026-10-28) -- without this, Brand
    would be silently blank for every row.
    """
    item_ids = []
    brand_by_id = {}
    page = 1
    while True:
        response = session.get(
            f"{api_domain}/inventory/v1/items",
            params={
                "organization_id": organization_id,
                "filter_by": "Status.Active",
                "page": page,
                "per_page": 200,
            },
            timeout=30,
        )
        response.raise_for_status()
        payload = response.json()
        for item in payload.get("items", []):
            item_ids.append(item["item_id"])
            if item.get("brand"):
                brand_by_id[item["item_id"]] = item["brand"]

        if not payload.get("page_context", {}).get("has_more_page"):
            break
        page += 1

    return item_ids, brand_by_id


def parse_retry_after(response, attempt):
    header_value = response.headers.get("Retry-After")
    if header_value is not None:
        try:
            return min(float(header_value), MAX_RETRY_AFTER_SECONDS)
        except ValueError:
            pass
    return min(2**attempt, MAX_RETRY_AFTER_SECONDS)


def chunked(sequence, size):
    for i in range(0, len(sequence), size):
        yield sequence[i : i + size]


def fetch_item_details_bulk(session, api_domain, organization_id, token_store, rate_limiter, item_ids):
    """Fetches full detail (including per-warehouse stock) for up to
    ITEMDETAILS_BATCH_SIZE items in a single call via /itemdetails.
    Returns None (after exhausting retries) rather than raising, so the
    caller can fall back to fetching this batch's items individually.
    """
    last_error = None
    for attempt in range(RETRY_ATTEMPTS):
        rate_limiter.wait()
        try:
            response = session.get(
                f"{api_domain}/inventory/v1/itemdetails",
                params={"organization_id": organization_id, "item_ids": ",".join(item_ids)},
                headers={"Authorization": f"Zoho-oauthtoken {token_store.get()}"},
                timeout=30,
            )
        except requests.exceptions.RequestException as exc:
            last_error = exc
            print(f"  [batch of {len(item_ids)}] attempt {attempt + 1} network error: {exc}", file=sys.stderr, flush=True)
            time.sleep(min(2**attempt, MAX_RETRY_AFTER_SECONDS))
            continue

        if response.status_code == 429:
            wait_s = parse_retry_after(response, attempt)
            last_error = requests.exceptions.HTTPError(f"429 Too Many Requests for batch of {len(item_ids)}")
            print(f"  [batch of {len(item_ids)}] attempt {attempt + 1} rate-limited (429), waiting {wait_s:.1f}s", file=sys.stderr, flush=True)
            time.sleep(wait_s)
            continue

        try:
            response.raise_for_status()
        except requests.exceptions.HTTPError as exc:
            last_error = exc
            print(f"  [batch of {len(item_ids)}] attempt {attempt + 1} HTTP error: {exc}", file=sys.stderr, flush=True)
            time.sleep(min(2**attempt, MAX_RETRY_AFTER_SECONDS))
            continue

        return response.json().get("items", [])

    print(f"  [batch of {len(item_ids)}] giving up after {RETRY_ATTEMPTS} attempts: {last_error}", file=sys.stderr, flush=True)
    return None


def fetch_item_detail(session, api_domain, organization_id, token_store, rate_limiter, item_id):
    last_error = None
    for attempt in range(RETRY_ATTEMPTS):
        rate_limiter.wait()
        try:
            response = session.get(
                f"{api_domain}/inventory/v1/items/{item_id}",
                params={"organization_id": organization_id},
                headers={"Authorization": f"Zoho-oauthtoken {token_store.get()}"},
                timeout=30,
            )
        except requests.exceptions.RequestException as exc:
            last_error = exc
            print(f"  [item {item_id}] attempt {attempt + 1} network error: {exc}", file=sys.stderr, flush=True)
            time.sleep(min(2**attempt, MAX_RETRY_AFTER_SECONDS))
            continue

        if response.status_code == 429:
            wait_s = parse_retry_after(response, attempt)
            last_error = requests.exceptions.HTTPError(f"429 Too Many Requests for item {item_id}")
            print(f"  [item {item_id}] attempt {attempt + 1} rate-limited (429), waiting {wait_s:.1f}s", file=sys.stderr, flush=True)
            time.sleep(wait_s)
            continue

        try:
            response.raise_for_status()
        except requests.exceptions.HTTPError as exc:
            last_error = exc
            print(f"  [item {item_id}] attempt {attempt + 1} HTTP error: {exc}", file=sys.stderr, flush=True)
            time.sleep(min(2**attempt, MAX_RETRY_AFTER_SECONDS))
            continue

        return response.json()["item"]

    print(f"  [item {item_id}] giving up after {RETRY_ATTEMPTS} attempts: {last_error}", file=sys.stderr, flush=True)
    return None


def fetch_active_items(session, api_domain, organization_id, token_store, rate_limiter):
    item_ids, brand_by_id = fetch_active_item_ids(session, api_domain, organization_id)
    batches = list(chunked(item_ids, ITEMDETAILS_BATCH_SIZE))
    print(
        f"Found {len(item_ids)} active item(s); fetching per-warehouse stock "
        f"detail in {len(batches)} batch(es) of up to {ITEMDETAILS_BATCH_SIZE}...",
        flush=True,
    )

    items = []
    completed = 0
    failed = 0
    stop_heartbeat = threading.Event()

    def heartbeat():
        while not stop_heartbeat.wait(HEARTBEAT_SECONDS):
            print(
                f"  ...still working: {completed}/{len(item_ids)} fetched ({failed} failed so far)",
                flush=True,
            )

    heartbeat_thread = threading.Thread(target=heartbeat, daemon=True)
    heartbeat_thread.start()

    try:
        with ThreadPoolExecutor(max_workers=DETAIL_FETCH_WORKERS) as executor:
            futures = {
                executor.submit(
                    fetch_item_details_bulk, session, api_domain, organization_id, token_store, rate_limiter, batch
                ): batch
                for batch in batches
            }
            for future in as_completed(futures):
                batch = futures[future]
                batch_items = future.result()
                if batch_items is None:
                    # The whole batch failed after retries -- fall back to
                    # per-item fetches so one bad id doesn't cost the rest.
                    print(f"  [batch of {len(batch)}] falling back to per-item fetch", file=sys.stderr, flush=True)
                    for item_id in batch:
                        detail = fetch_item_detail(
                            session, api_domain, organization_id, token_store, rate_limiter, item_id
                        )
                        completed += 1
                        if detail is None:
                            failed += 1
                        else:
                            items.append(detail)
                else:
                    completed += len(batch)
                    missing = len(batch) - len(batch_items)
                    if missing:
                        failed += missing
                        print(
                            f"  [batch of {len(batch)}] returned {len(batch_items)} of {len(batch)} items",
                            file=sys.stderr,
                            flush=True,
                        )
                    items.extend(batch_items)
                print(f"  ...{completed}/{len(item_ids)} items fetched ({failed} failed)", flush=True)
    finally:
        stop_heartbeat.set()
        heartbeat_thread.join()

    if failed:
        print(f"Warning: {failed} item(s) could not be fetched after retries and were skipped.", file=sys.stderr, flush=True)

    # The bulk /itemdetails endpoint doesn't return `brand`, unlike the
    # single-item endpoint and the items list -- patch it back in from the
    # list call above (see fetch_active_item_ids docstring).
    for item in items:
        if not item.get("brand"):
            brand = brand_by_id.get(item.get("item_id"))
            if brand:
                item["brand"] = brand

    return items


def find_custom_field(item, *label_fragments):
    for field in item.get("custom_fields") or []:
        label = (field.get("label") or "").strip().lower()
        if any(fragment in label for fragment in label_fragments):
            value = field.get("value")
            if value:
                return value
    return ""


def extract_barcode(item):
    for field_name in BARCODE_FIELDS:
        value = item.get(field_name)
        if value:
            return value
    return find_custom_field(item, "bar code", "barcode")


def extract_brand(item):
    return item.get("brand") or find_custom_field(item, "brand")


def build_rows(items, warehouse_name):
    warehouse_name = warehouse_name.strip().lower()
    rows = []
    for item in items:
        for warehouse in item.get("warehouses") or []:
            if (warehouse.get("warehouse_name") or "").strip().lower() != warehouse_name:
                continue
            rows.append(
                {
                    "BAR CODE": extract_barcode(item),
                    "SKU": item.get("sku", ""),
                    "Item Name": item.get("name", ""),
                    "SOH": warehouse.get("warehouse_stock_on_hand", ""),
                    "Sales Price": item.get("rate", ""),
                    "Brand": extract_brand(item),
                }
            )
            break
    rows.sort(key=lambda row: (row["SKU"] or "").strip().lower())
    return rows


def write_excel(rows, output_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(OUTPUT_COLUMNS)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"

    for row in rows:
        sheet.append([row[column] for column in OUTPUT_COLUMNS])

    soh_col = OUTPUT_COLUMNS.index("SOH") + 1
    price_col = OUTPUT_COLUMNS.index("Sales Price") + 1
    for row_idx in range(2, len(rows) + 2):
        sheet.cell(row=row_idx, column=soh_col).number_format = "#,##0"
        sheet.cell(row=row_idx, column=price_col).number_format = "#,##0.00"

    for col_idx, column_name in enumerate(OUTPUT_COLUMNS, start=1):
        longest = max(
            [len(column_name)] + [len(str(row[column_name])) for row in rows],
            default=len(column_name),
        )
        sheet.column_dimensions[get_column_letter(col_idx)].width = longest + 4

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook.save(output_path)


def main():
    warehouse_name = os.environ.get("WAREHOUSE_NAME", "Store 1")
    report_tz = ZoneInfo(os.environ.get("REPORT_TZ", "Africa/Nairobi"))

    report_date = datetime.now(report_tz)
    output_path = f"reports/SOH_{report_date.strftime('%d_%m_%Y')}.xlsx"

    rate_limit_per_minute = int(os.environ.get("RATE_LIMIT_PER_MINUTE", DEFAULT_RATE_LIMIT_PER_MINUTE))
    rate_limiter = RateLimiter(rate_limit_per_minute)

    access_token, api_domain = get_access_token()
    token_store = TokenStore(access_token)
    session = requests.Session()
    session.headers["Authorization"] = f"Zoho-oauthtoken {access_token}"

    organization_id = resolve_organization_id(session, api_domain)
    items = fetch_active_items(session, api_domain, organization_id, token_store, rate_limiter)
    rows = build_rows(items, warehouse_name)

    if not rows:
        seen_warehouses = sorted(
            {
                warehouse.get("warehouse_name", "")
                for item in items
                for warehouse in (item.get("warehouses") or [])
            }
        )
        print(
            f"Warning: no active items found stocked at warehouse '{warehouse_name}'. "
            f"Checked {len(items)} active item(s); warehouse names seen: {seen_warehouses}",
            file=sys.stderr,
        )

    write_excel(rows, output_path)
    print(f"Wrote {len(rows)} rows to {output_path}")

    github_output = os.environ.get("GITHUB_OUTPUT")
    if github_output:
        with open(github_output, "a") as f:
            f.write(f"report_path={output_path}\n")


if __name__ == "__main__":
    main()
