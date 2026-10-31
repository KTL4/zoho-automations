import datetime
import json
import os
import sys
import tempfile
import time
import unittest
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import container_sellout_report as report  # noqa: E402

D = datetime.date
PATTERNS = report.supplier_patterns()


def lot(seq, day, qty_in, on_hand, sku="A1", number="C1", contact="Uncle Bill Ltd", ttype="Bill", cost=10.0):
    return {"seq": seq, "date": day, "transaction_type": ttype, "transaction_number": number,
            "contact_name": contact, "lot_number": "", "sku": sku, "item_name": f"Item {sku}",
            "quantity_in": float(qty_in), "quantity_on_hand": float(on_hand), "unit_cost": cost}


class AllocateSalesFifoTests(unittest.TestCase):
    def test_oldest_lot_fills_first_and_room_is_capped_by_zoho_consumed(self):
        lots = [lot(0, D(2026, 1, 1), 100, 40), lot(1, D(2026, 2, 1), 100, 100)]
        allocations, unmatched = report.allocate_sales_fifo(lots, {D(2026, 2, 10): 70})
        # Lot 0 only had 60 consumed per Zoho; lot 1 shows nothing consumed, so 10 is unmatched.
        self.assertEqual(allocations[0], [(D(2026, 2, 10), 60)])
        self.assertNotIn(1, allocations)
        self.assertEqual(unmatched, 10)

    def test_sale_never_goes_to_a_lot_that_had_not_arrived(self):
        lots = [lot(0, D(2026, 3, 1), 50, 0)]
        allocations, unmatched = report.allocate_sales_fifo(lots, {D(2026, 2, 1): 5, D(2026, 3, 2): 5})
        self.assertEqual(allocations[0], [(D(2026, 3, 2), 5)])
        self.assertEqual(unmatched, 5)

    def test_net_return_is_carried_into_next_sale(self):
        lots = [lot(0, D(2026, 1, 1), 50, 0)]
        allocations, _ = report.allocate_sales_fifo(lots, {D(2026, 1, 5): -3, D(2026, 1, 6): 10})
        self.assertEqual(allocations[0], [(D(2026, 1, 6), 7)])


class SummaryTests(unittest.TestCase):
    def test_container_metrics(self):
        lots = [lot(0, D(2026, 1, 1), 100, 10), lot(1, D(2026, 1, 1), 100, 100, sku="B2")]
        daily = {"A1": {D(2026, 1, 11): 50, D(2026, 2, 20): 30}}
        rows, sku_rows, unmatched = report.build_report_data(lots, daily, PATTERNS, D(2026, 4, 1))
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual((r["supplier"], r["container"], r["skus"]), ("UBL", "C1", 2))
        self.assertEqual(r["units_in"], 200)
        self.assertAlmostEqual(r["sellout"], 90 / 200)
        self.assertEqual(r["sold"], 80)
        self.assertEqual(r["other_out"], 10)
        self.assertAlmostEqual(r["sold_by"][30], 50 / 200)
        self.assertAlmostEqual(r["sold_by"][60], 80 / 200)
        self.assertIsNone(r["sold_by"][180])  # container only 90 days old
        self.assertIsNone(r["days_to"][0.5])   # 40% sold, never reached 50%
        a1 = next(s for s in sku_rows if s["sku"] == "A1")
        self.assertEqual(a1["days_to"][0.5], 10)
        self.assertEqual(a1["days_to"][0.8], 50)
        self.assertEqual(unmatched, 0)

    def test_non_container_lots_are_fifo_competitors_but_not_reported(self):
        lots = [lot(0, D(2025, 12, 1), 20, 0, ttype="Opening Stock", contact=""),
                lot(1, D(2026, 1, 1), 100, 70)]
        rows, _, _ = report.build_report_data(lots, {"A1": {D(2026, 1, 5): 50}}, PATTERNS, D(2026, 4, 1))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["sold"], 30)  # first 20 came out of opening stock

    def test_supplier_grouping(self):
        self.assertEqual(report.supplier_of("UNCLE BILL LIMITED", PATTERNS), "UBL")
        self.assertEqual(report.supplier_of("Aiwibi Trading", PATTERNS), "Aiwibi")
        self.assertEqual(report.supplier_of("EFRIS STOCK", PATTERNS), "EFRIS STOCK")
        # Aiwibi containers A04 onward are billed to this vendor; the first
        # live report missed all 17 of them.
        self.assertEqual(report.supplier_of("AWB Inventory Receipt (USD)", PATTERNS), "Aiwibi")
        self.assertIsNone(report.supplier_of("Local Hardware", PATTERNS))


def fake_response(payload, status=200):
    response = MagicMock()
    response.status_code = status
    response.json.return_value = payload
    response.headers = {}
    response.raise_for_status = MagicMock()
    return response


class IncrementalRefreshTests(unittest.TestCase):
    def test_merge_takes_old_lots_from_cache_and_recent_from_fetch(self):
        cached = [lot(0, D(2024, 1, 1), 10, 5), lot(1, D(2025, 6, 1), 10, 9)]
        fetched = [lot(0, D(2024, 1, 1), 10, 0),   # pre-cutoff row in the pull is ignored
                   lot(1, D(2025, 6, 1), 10, 2), lot(2, D(2025, 7, 1), 10, 2)]
        merged = report.merge_lots(cached, fetched, D(2025, 1, 1))
        self.assertEqual([(l["date"], l["quantity_on_hand"]) for l in merged],
                         [(D(2024, 1, 1), 5), (D(2025, 6, 1), 2), (D(2025, 7, 1), 2)])
        self.assertEqual([l["seq"] for l in merged], [0, 1, 2])

    def test_repeated_lot_rows_are_counted_once(self):
        # Live data, 2026-10-31: lot 2871458 (LCL W2, 2,400 wet wipes) came
        # back twice; Zoho's bill shows it once.
        row = {"date": "2025-12-13", "transaction_type": "bill", "transaction_number": "LCL W2 Inventory",
               "contact_name": "AWB Inventory Receipt (USD)", "lot_number": "2871458", "sku": "AWB80-03",
               "quantity_in": "2400", "quantity_on_hand": "0", "cost_price_per_unit": "1"}
        other = {**row, "lot_number": "2871459", "sku": "AWB80-04"}
        page = {"fifo_cost_lot_tracking": [row, row, other, row], "page_context": {"has_more_page": False}}
        session = MagicMock()
        session.get.return_value = fake_response(page)
        token_store = MagicMock()
        token_store.get.return_value = "tok"
        with patch.object(report.time, "sleep"):
            lots = report.fetch_lots(session, "https://api", "1", token_store, D(2025, 1, 1), D(2026, 1, 1))
        self.assertEqual([(l["lot_number"], l["quantity_in"]) for l in lots], [("2871458", 2400), ("2871459", 2400)])
        self.assertEqual([l["seq"] for l in lots], [0, 1])

    def test_saved_lots_with_repeats_load_once(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(report, "DATA_DIR", tmp):
            one = {**lot(0, D(2025, 12, 13), 2400, 0), "lot_number": "2871458"}
            report.save_lots("42", [one, {**one, "seq": 1}, {**one, "seq": 2, "lot_number": ""},
                                    {**one, "seq": 3, "lot_number": ""}])
            loaded = report.load_lots("42")
        self.assertEqual([l["lot_number"] for l in loaded], ["2871458", "", ""])  # blank lot numbers aren't merged

    def test_plan_refresh(self):
        today = D(2026, 10, 30)
        with patch.dict(os.environ, {"FULL_REFRESH": "", "FULL_REFRESH_DAYS": ""}):
            self.assertTrue(report.plan_refresh({}, today, True)[0])
            self.assertTrue(report.plan_refresh({"last_full_refresh": "2026-10-29"}, today, False)[0])
            self.assertFalse(report.plan_refresh({"last_full_refresh": "2026-10-20"}, today, True)[0])
            self.assertTrue(report.plan_refresh({"last_full_refresh": "2026-10-02"}, today, True)[0])
        with patch.dict(os.environ, {"FULL_REFRESH": "true"}):
            self.assertTrue(report.plan_refresh({"last_full_refresh": "2026-10-29"}, today, True)[0])

    def test_windowed_sync_keeps_older_cached_invoices(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(report, "DATA_DIR", tmp):
            report.save_cache(report.cache_path("42"), {
                ("invoice", "old"): {"date": "2025-01-05", "modified": "m1", "lines": [("A1", "", 3.0)]},
                ("invoice", "new"): {"date": "2026-10-01", "modified": "m1", "lines": [("A1", "", 2.0)]},
            })
            listed = {"new": ("2026-10-01", "m1")}  # old invoice falls outside the listed window
            with patch.object(report, "list_documents", return_value=listed):
                cache, complete = report.sync_sales(MagicMock(), "https://api", "42", MagicMock(), MagicMock(),
                                                    D(2026, 8, 1), time.monotonic() + 60)
        self.assertTrue(complete)
        self.assertEqual(set(cache), {("invoice", "old"), ("invoice", "new")})


class TokenHandlingTests(unittest.TestCase):
    def test_401_refreshes_token_once_and_retries(self):
        session = MagicMock()
        session.get.side_effect = [fake_response({}, status=401), fake_response({"ok": 1})]
        token_store = MagicMock()
        token_store.get.side_effect = ["old", "new"]
        response = report.authorized_get(session, token_store, "https://api/x/invoices", {}, timeout=5)
        self.assertEqual(response.status_code, 200)
        token_store.force_refresh.assert_called_once()
        self.assertEqual(session.get.call_args.kwargs["headers"]["Authorization"], "Zoho-oauthtoken new")

    def test_in_progress_report_waits_without_minting_tokens(self):
        # The first live run minted a token on every "still computing" retry,
        # which got the working token revoked.
        busy = {"fifo_cost_lot_tracking": [], "page_context": {"is_inv_txn_job_in_progress": True, "has_more_page": True}}
        ready = {"fifo_cost_lot_tracking": [{"quantity_in": 1}], "page_context": {"has_more_page": True}}
        session = MagicMock()
        session.get.side_effect = [fake_response(busy), fake_response(ready)]
        token_store = MagicMock()
        token_store.get.return_value = "tok"
        with patch.object(report.time, "sleep"):
            rows, more = report.fetch_lot_page(session, "https://api", "1", token_store, D(2020, 1, 1), D(2026, 1, 1), 242)
        self.assertEqual(len(rows), 1)
        self.assertTrue(more)
        token_store.force_refresh.assert_not_called()


class EndToEndTests(unittest.TestCase):
    def test_full_run_writes_report_and_cache_then_reuses_cache(self):
        lot_rows = [
            {"date": "2026-01-01", "transaction_type": "Bill", "transaction_number": "C30", "contact_name": "Uncle Bill",
             "sku": "A1", "item_name": "Widget", "quantity_in": "1,000", "quantity_on_hand": "400", "cost_price_per_unit": "5"},
            {"date": "2026-01-01", "transaction_type": "Invoice", "transaction_number": "INV-1", "contact_name": "Shop",
             "sku": "A1", "quantity_in": 0, "quantity_on_hand": 0},
        ] + [{"date": "2025-01-01", "transaction_type": "Opening Stock", "transaction_number": f"OS{i}",
              "contact_name": "", "sku": f"Z{i}", "quantity_in": 1, "quantity_on_hand": 1} for i in range(60)]
        invoices = [{"invoice_id": "i1", "date": "2026-01-10", "status": "paid", "last_modified_time": "t1"},
                    {"invoice_id": "i2", "date": "2026-01-12", "status": "void", "last_modified_time": "t1"}]
        detail_calls = []

        lot_pulls = []

        def get(url, params=None, **_):
            if url.endswith("/reports/fifocostlottracking"):
                lot_pulls.append(params["from_date"])
                return fake_response({"fifo_cost_lot_tracking": lot_rows, "page_context": {"has_more_page": False}})
            if url.endswith("/invoices"):
                return fake_response({"invoices": invoices, "page_context": {"has_more_page": False}})
            if url.endswith("/creditnotes"):
                return fake_response({"creditnotes": [], "page_context": {"has_more_page": False}})
            raise AssertionError(url)

        def fetch_detail(session, api_domain, org, token_store, limiter, url, result_key, label):
            detail_calls.append(url)
            return {"line_items": [{"sku": "a1", "quantity": 600}]}

        session = MagicMock()
        session.get.side_effect = get
        session.headers = {}
        with tempfile.TemporaryDirectory() as tmp, \
                patch.object(report.requests, "Session", return_value=session), \
                patch.object(report.common, "get_access_token", return_value=("tok", "https://api")), \
                patch.object(report.common, "fetch_detail", side_effect=fetch_detail), \
                patch.object(report.time, "sleep"), \
                patch.dict(os.environ, {"ZOHO_ORGANIZATION_ID": "42", "REPORT_AS_OF_DATE": "2026-04-01"}):
            cwd = os.getcwd()
            os.chdir(tmp)
            try:
                report.main()
                report.main()  # second run: incremental, and the invoice is cached
                with open(os.path.join("data", "sellout_state_42.json")) as handle:
                    state = json.load(handle)
                wb = load_workbook(os.path.join("reports", "Container_Sellout_01_04_2026.xlsx"))
                self.assertTrue(os.path.exists(os.path.join("data", "sales_lines_42.csv")))
                self.assertTrue(os.path.exists(os.path.join("data", "fifo_lots_42.csv")))
            finally:
                os.chdir(cwd)

        self.assertEqual(len(detail_calls), 1)
        self.assertEqual(lot_pulls, ["2015-01-01", "2025-04-01"])  # full, then only the last 365 days
        self.assertEqual(state, {"last_full_refresh": "2026-04-01", "last_run": "2026-04-01",
                                 "sales_backfill_complete": True})
        sheet = wb["Containers"]
        headers = [c.value for c in sheet[1]]
        row = dict(zip(headers, [c.value for c in sheet[2]]))
        self.assertEqual(row["Container"], "C30")
        self.assertEqual(row["Units In"], 1000)
        self.assertAlmostEqual(row["Sellout % (Zoho FIFO)"], 0.6)
        self.assertEqual(row["Units Sold"], 600)
        self.assertAlmostEqual(row["Sold % by Day 30"], 0.6)
        self.assertEqual(row["Days to 50% Sold"], 9)
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(wb.sheetnames, ["Containers", "Container x SKU", "Read Me"])


if __name__ == "__main__":
    unittest.main()
