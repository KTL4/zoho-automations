import os
import shutil
import sys
import unittest
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import item_runway_report as report  # noqa: E402
import sales_by_item_common as common  # noqa: E402


class MainIntegrationTest(unittest.TestCase):
    """End-to-end wiring check with every HTTP call mocked.

    Verifies this script's own specific behavior (14-period window anchored
    on the last full week *before* "today", month-name filename); the
    shared fetch/aggregate/filter/format logic itself is covered in
    tests/test_sales_by_item_common.py.
    """

    @patch.dict(
        os.environ,
        {
            "ZOHO_CLIENT_ID": "id",
            "ZOHO_CLIENT_SECRET": "secret",
            "ZOHO_REFRESH_TOKEN": "refresh",
            "ZOHO_ORGANIZATION_ID": "org1",
            "REPORT_TZ": "UTC",
            "RATE_LIMIT_PER_MINUTE": "6000",
            # Thursday -- that week's Monday is 2026-09-28, so "previous
            # week" (the most recent full week strictly before today's own
            # week) is 2026-09-21 .. 2026-09-27, entirely in September.
            "REPORT_AS_OF_DATE": "2026-10-01",
        },
        clear=True,
    )
    @patch("sales_by_item_common.requests.Session")
    @patch("sales_by_item_common.requests.post")
    def test_main_writes_expected_file(self, mock_post, mock_session_cls):
        mock_post.return_value = MagicMock(
            raise_for_status=MagicMock(),
            json=MagicMock(return_value={"access_token": "tok", "api_domain": "https://api"}),
        )

        session = MagicMock()
        mock_session_cls.return_value = session

        items_page = MagicMock()
        items_page.raise_for_status = MagicMock()
        items_page.json.return_value = {
            "items": [{"item_id": "item-1"}, {"item_id": "item-2"}, {"item_id": "item-3"}],
            "page_context": {"has_more_page": False},
        }
        item_detail = MagicMock()
        item_detail.status_code = 200
        item_detail.raise_for_status = MagicMock()
        item_detail.json.return_value = {
            "item": {
                "item_id": "item-1",
                "sku": "SKU1",
                "name": "Widget",
                "brand": "Acme",
                "item_type": "inventory",
            }
        }
        # Non-stock catalog entry (e.g. a freight/container line item) -- should
        # never appear in the output regardless of whether it has sales.
        container_detail = MagicMock()
        container_detail.status_code = 200
        container_detail.raise_for_status = MagicMock()
        container_detail.json.return_value = {
            "item": {
                "item_id": "item-2",
                "sku": "N/A",
                "name": "20FT Container",
                "item_type": "purchases",
            }
        }
        # Real stock item with zero sales in the window -- should be omitted.
        no_sales_detail = MagicMock()
        no_sales_detail.status_code = 200
        no_sales_detail.raise_for_status = MagicMock()
        no_sales_detail.json.return_value = {
            "item": {
                "item_id": "item-3",
                "sku": "SKU3",
                "name": "Unsold Widget",
                "item_type": "inventory",
            }
        }

        invoices_page = MagicMock()
        invoices_page.raise_for_status = MagicMock()
        invoices_page.json.return_value = {
            "invoices": [{"invoice_id": "inv-1", "date": "2026-09-22", "status": "paid"}],
            "page_context": {"has_more_page": False},
        }
        invoice_detail = MagicMock()
        invoice_detail.status_code = 200
        invoice_detail.raise_for_status = MagicMock()
        invoice_detail.json.return_value = {
            "invoice": {
                "invoice_id": "inv-1",
                "date": "2026-09-22",
                "line_items": [{"item_id": "item-1", "quantity": 4}],
            }
        }

        def get_side_effect(url, params=None, headers=None, timeout=None):
            if url.endswith("/inventory/v1/items"):
                return items_page
            if url.endswith("/inventory/v1/items/item-1"):
                return item_detail
            if url.endswith("/inventory/v1/items/item-2"):
                return container_detail
            if url.endswith("/inventory/v1/items/item-3"):
                return no_sales_detail
            if url.endswith("/inventory/v1/invoices"):
                return invoices_page
            if url.endswith("/inventory/v1/invoices/inv-1"):
                return invoice_detail
            raise AssertionError(f"Unexpected GET {url}")

        session.get.side_effect = get_side_effect

        # This is the real production output path -- a live run may have
        # already committed a real report there, so back it up rather than
        # deleting it outright (a prior version of this test did exactly
        # that and destroyed the real committed file when run locally).
        output_path = "reports/Sales_by_item(sep).xlsx"
        backup_path = output_path + ".bak"
        preexisting = os.path.exists(output_path)
        if preexisting:
            shutil.move(output_path, backup_path)

        try:
            report.main()
            self.assertTrue(os.path.exists(output_path))

            import openpyxl

            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook.active
            self.assertEqual(sheet.cell(row=1, column=1).value, "SKU")
            self.assertEqual([c.value for c in sheet[1]][:3], ["SKU", "Item Name", "Brand"])
            self.assertEqual(len(list(sheet[1])), 3 + 14)  # 14-period window
            # Only item-1 survives: item-2 is a non-stock ("purchases") catalog
            # entry and item-3 has zero sales in the window.
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet.cell(row=2, column=1).value, "SKU1")
            self.assertEqual(sheet.cell(row=2, column=2).value, "Widget")
            self.assertEqual(sheet.cell(row=2, column=3).value, "Acme")
            self.assertEqual(sheet.cell(row=2, column=sheet.max_column).value, 4)
        finally:
            if preexisting:
                shutil.move(backup_path, output_path)
            elif os.path.exists(output_path):
                os.remove(output_path)


if __name__ == "__main__":
    unittest.main()
