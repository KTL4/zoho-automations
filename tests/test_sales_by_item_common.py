import datetime
import os
import sys
import unittest
from unittest.mock import MagicMock

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import sales_by_item_common as common  # noqa: E402


class ComputeWeekBucketsEndingByTests(unittest.TestCase):
    def test_reference_date_on_sunday_includes_that_weeks_own_end(self):
        # 2026-09-13 is itself a Sunday.
        buckets = common.compute_week_buckets_ending_by(datetime.date(2026, 9, 13), 1)
        self.assertEqual(buckets, [(datetime.date(2026, 9, 7), datetime.date(2026, 9, 13))])

    def test_reference_date_midweek_uses_previous_sunday(self):
        # 2026-09-16 is a Wednesday, inside the still-open Sep14-20 week.
        buckets = common.compute_week_buckets_ending_by(datetime.date(2026, 9, 16), 1)
        self.assertEqual(buckets, [(datetime.date(2026, 9, 7), datetime.date(2026, 9, 13))])

    def test_reference_date_on_monday_uses_previous_sunday(self):
        buckets = common.compute_week_buckets_ending_by(datetime.date(2026, 9, 14), 1)
        self.assertEqual(buckets, [(datetime.date(2026, 9, 7), datetime.date(2026, 9, 13))])

    def test_multiple_periods_are_chronological_and_contiguous(self):
        buckets = common.compute_week_buckets_ending_by(datetime.date(2026, 9, 13), 36)

        self.assertEqual(len(buckets), 36)
        self.assertEqual(buckets[-1], (datetime.date(2026, 9, 7), datetime.date(2026, 9, 13)))
        for (start, end) in buckets:
            self.assertEqual((end - start).days, 6)
        for i in range(1, len(buckets)):
            self.assertEqual(buckets[i][0] - buckets[i - 1][0], datetime.timedelta(days=7))

    def test_matches_monthly_reports_previous_full_week_semantics(self):
        # The monthly report anchors on (today - 1 day) to get "the last full
        # week strictly before today's week" -- verify against the values
        # already validated by two live production runs.
        today = datetime.date(2026, 10, 1)
        buckets = common.compute_week_buckets_ending_by(today - datetime.timedelta(days=1), 14)

        self.assertEqual(len(buckets), 14)
        self.assertEqual(buckets[-1], (datetime.date(2026, 9, 21), datetime.date(2026, 9, 27)))
        self.assertEqual(buckets[0], (datetime.date(2026, 6, 22), datetime.date(2026, 6, 28)))


class WeekLabelTests(unittest.TestCase):
    def test_label_uses_iso_week_number(self):
        self.assertEqual(common.week_label(datetime.date(2026, 9, 21)), "WK 39")


class AggregateSalesTests(unittest.TestCase):
    def setUp(self):
        self.buckets = common.compute_week_buckets_ending_by(datetime.date(2026, 9, 27), 14)

    def test_sums_quantities_into_correct_week_bucket(self):
        invoices = [
            {
                "date": "2026-09-22",  # falls in the most recent bucket
                "line_items": [
                    {"item_id": "item-1", "quantity": 3},
                    {"item_id": "item-2", "quantity": 1},
                ],
            },
            {
                "date": "2026-09-25",  # same week as above
                "line_items": [{"item_id": "item-1", "quantity": 2}],
            },
            {
                "date": "2026-06-24",  # oldest bucket (2026-06-22 .. 2026-06-28)
                "line_items": [{"item_id": "item-1", "quantity": 10}],
            },
        ]
        sales = common.aggregate_sales(invoices, self.buckets)

        self.assertEqual(sales["item-1"][-1], 5)
        self.assertEqual(sales["item-2"][-1], 1)
        self.assertEqual(sales["item-1"][0], 10)
        self.assertEqual(sum(sales["item-2"]), 1)

    def test_skips_invoices_outside_window_and_missing_dates(self):
        invoices = [
            {"date": "2020-01-01", "line_items": [{"item_id": "item-1", "quantity": 99}]},
            {"line_items": [{"item_id": "item-1", "quantity": 5}]},  # no date
            None,
        ]
        sales = common.aggregate_sales(invoices, self.buckets)
        self.assertEqual(sales, {})

    def test_line_items_without_item_id_are_ignored(self):
        invoices = [{"date": "2026-09-22", "line_items": [{"quantity": 5}]}]
        sales = common.aggregate_sales(invoices, self.buckets)
        self.assertEqual(sales, {})


class FetchInvoiceIdsInWindowTests(unittest.TestCase):
    def _make_response(self, invoices, has_more_page):
        response = MagicMock()
        response.raise_for_status = MagicMock()
        response.json.return_value = {
            "invoices": invoices,
            "page_context": {"has_more_page": has_more_page},
        }
        return response

    def test_paginates_and_stops_past_window_start(self):
        window_start = datetime.date(2026, 9, 1)
        window_end = datetime.date(2026, 9, 30)

        page1 = self._make_response(
            [
                {"invoice_id": "future", "date": "2026-10-05", "status": "paid"},  # newer than window
                {"invoice_id": "inv-1", "date": "2026-09-20", "status": "paid"},
                {"invoice_id": "inv-draft", "date": "2026-09-15", "status": "draft"},
                {"invoice_id": "inv-void", "date": "2026-09-10", "status": "void"},
            ],
            has_more_page=True,
        )
        page2 = self._make_response(
            [
                {"invoice_id": "inv-2", "date": "2026-09-05", "status": "sent"},
                {"invoice_id": "too-old", "date": "2026-08-20", "status": "paid"},  # should stop here
                {"invoice_id": "never-fetched", "date": "2026-08-01", "status": "paid"},
            ],
            has_more_page=True,  # would keep going if not for the window_start cutoff
        )

        session = MagicMock()
        session.get.side_effect = [page1, page2]

        ids = common.fetch_invoice_ids_in_window(session, "https://api", "org1", window_start, window_end)

        self.assertEqual(ids, ["inv-1", "inv-2"])
        self.assertEqual(session.get.call_count, 2)

    def test_stops_when_no_more_pages(self):
        window_start = datetime.date(2026, 9, 1)
        window_end = datetime.date(2026, 9, 30)
        page1 = self._make_response(
            [{"invoice_id": "inv-1", "date": "2026-09-15", "status": "paid"}],
            has_more_page=False,
        )
        session = MagicMock()
        session.get.side_effect = [page1]

        ids = common.fetch_invoice_ids_in_window(session, "https://api", "org1", window_start, window_end)
        self.assertEqual(ids, ["inv-1"])
        self.assertEqual(session.get.call_count, 1)


class ExtractBrandTests(unittest.TestCase):
    def test_prefers_top_level_brand_field(self):
        item = {"brand": "Acme", "custom_fields": [{"label": "Brand", "value": "Other"}]}
        self.assertEqual(common.extract_brand(item), "Acme")

    def test_falls_back_to_custom_field(self):
        item = {"custom_fields": [{"label": "Brand Name", "value": "Acme"}]}
        self.assertEqual(common.extract_brand(item), "Acme")

    def test_missing_brand_returns_empty_string(self):
        self.assertEqual(common.extract_brand({}), "")


class IsStockItemTests(unittest.TestCase):
    def test_inventory_type_is_kept(self):
        self.assertTrue(common.is_stock_item({"item_type": "inventory"}))

    def test_non_inventory_types_are_excluded(self):
        for item_type in ["purchases", "sales", "sales_and_purchases", "service"]:
            with self.subTest(item_type=item_type):
                self.assertFalse(common.is_stock_item({"item_type": item_type}))

    def test_case_insensitive(self):
        self.assertTrue(common.is_stock_item({"item_type": "Inventory"}))

    def test_missing_item_type_fails_open(self):
        self.assertTrue(common.is_stock_item({}))
        self.assertTrue(common.is_stock_item({"item_type": ""}))


class BuildItemCatalogTests(unittest.TestCase):
    def test_excludes_non_stock_items_from_catalog(self):
        session = MagicMock()

        items_page = MagicMock()
        items_page.raise_for_status = MagicMock()
        items_page.json.return_value = {
            "items": [{"item_id": "item-1"}, {"item_id": "item-2"}],
            "page_context": {"has_more_page": False},
        }

        stock_item_detail = MagicMock()
        stock_item_detail.raise_for_status = MagicMock()
        stock_item_detail.json.return_value = {
            "item": {"item_id": "item-1", "sku": "SKU1", "name": "Widget", "item_type": "inventory"}
        }
        container_detail = MagicMock()
        container_detail.raise_for_status = MagicMock()
        container_detail.json.return_value = {
            "item": {"item_id": "item-2", "sku": "N/A", "name": "20FT Container", "item_type": "purchases"}
        }

        def get_side_effect(url, params=None, headers=None, timeout=None):
            if url.endswith("/inventory/v1/items"):
                return items_page
            if url.endswith("/inventory/v1/items/item-1"):
                return stock_item_detail
            if url.endswith("/inventory/v1/items/item-2"):
                return container_detail
            raise AssertionError(f"Unexpected GET {url}")

        session.get.side_effect = get_side_effect

        token_store = MagicMock()
        token_store.get.return_value = "tok"
        rate_limiter = MagicMock()

        catalog = common.build_item_catalog(session, "https://api", "org1", token_store, rate_limiter)

        self.assertEqual(list(catalog.keys()), ["item-1"])
        self.assertEqual(catalog["item-1"]["Item Name"], "Widget")


class BuildRowsTests(unittest.TestCase):
    def test_excludes_zero_sales_items_and_sorts_remaining_by_name(self):
        catalog = {
            "item-1": {"SKU": "SKU1", "Item Name": "Zebra Widget", "Brand": "Acme"},
            "item-2": {"SKU": "SKU2", "Item Name": "Apple Widget", "Brand": "Acme"},
            "item-3": {"SKU": "SKU3", "Item Name": "No Sales Widget", "Brand": "Acme"},
        }
        buckets = [(datetime.date(2026, 1, 1), datetime.date(2026, 1, 7))]
        sales = {"item-1": [5], "item-2": [0]}

        rows = common.build_rows(catalog, sales, buckets)

        # item-2 (all-zero sales) and item-3 (no sales at all) are both omitted.
        self.assertEqual([row[1] for row in rows], ["Zebra Widget"])
        self.assertEqual(rows[0][3], 5)

    def test_item_with_any_nonzero_week_is_kept(self):
        catalog = {"item-1": {"SKU": "SKU1", "Item Name": "Widget", "Brand": "Acme"}}
        buckets = [
            (datetime.date(2026, 1, 1), datetime.date(2026, 1, 7)),
            (datetime.date(2026, 1, 8), datetime.date(2026, 1, 14)),
        ]
        sales = {"item-1": [0, 3]}

        rows = common.build_rows(catalog, sales, buckets)

        self.assertEqual(len(rows), 1)

    def test_sales_for_items_outside_catalog_are_dropped_with_warning(self):
        catalog = {"item-1": {"SKU": "SKU1", "Item Name": "Widget", "Brand": "Acme"}}
        buckets = [(datetime.date(2026, 1, 1), datetime.date(2026, 1, 7))]
        sales = {"item-1": [5], "discontinued-item": [10]}

        rows = common.build_rows(catalog, sales, buckets)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][1], "Widget")


class WriteExcelTests(unittest.TestCase):
    def test_writes_header_frozen_panes_and_number_format(self):
        import openpyxl
        from unittest.mock import patch

        rows = [
            ["SKU1", "Widget A", "Acme", 3, 0],
            ["SKU2", "Widget B", "Acme", 0, 7],
        ]
        week_columns = ["WK 1", "WK 2"]

        with patch("os.makedirs"):
            output_path = os.path.join(
                os.path.dirname(__file__), "_tmp_test_output.xlsx"
            )
            common.write_excel(rows, week_columns, output_path)

        try:
            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook.active

            self.assertEqual(
                [cell.value for cell in sheet[1]],
                ["SKU", "Item Name", "Brand", "WK 1", "WK 2"],
            )
            self.assertTrue(sheet[1][0].font.bold)
            self.assertEqual(sheet.freeze_panes, "B2")
            self.assertEqual(sheet.cell(row=2, column=4).number_format, "#,##0")
            self.assertEqual(sheet.cell(row=2, column=1).value, "SKU1")
        finally:
            if os.path.exists(output_path):
                os.remove(output_path)


class ResolveTodayTests(unittest.TestCase):
    def test_uses_override_when_set(self):
        from unittest.mock import patch
        from zoneinfo import ZoneInfo

        with patch.dict(os.environ, {"REPORT_AS_OF_DATE": "2026-09-13"}, clear=False):
            today = common.resolve_today(ZoneInfo("UTC"))
        self.assertEqual(today, datetime.date(2026, 9, 13))

    def test_uses_real_now_when_unset(self):
        from unittest.mock import patch
        from zoneinfo import ZoneInfo

        with patch.dict(os.environ, {}, clear=False):
            os.environ.pop("REPORT_AS_OF_DATE", None)
            today = common.resolve_today(ZoneInfo("UTC"))
        self.assertEqual(today, datetime.datetime.now(ZoneInfo("UTC")).date())


if __name__ == "__main__":
    unittest.main()
