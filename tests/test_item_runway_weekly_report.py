import os
import shutil
import sys
import unittest
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import item_runway_weekly_report as report  # noqa: E402


def _mock_response(json_value):
    response = MagicMock()
    response.raise_for_status = MagicMock()
    response.json.return_value = json_value
    return response


class MainIntegrationTest(unittest.TestCase):
    """End-to-end wiring check with every HTTP call mocked.

    Verifies this script's own specific behavior (36-period window anchored
    on "today" itself, so a Sunday run includes the week ending that day,
    and the WKx-WKy filename pattern); the shared fetch/aggregate/filter/
    format logic itself is covered in tests/test_sales_by_item_common.py.
    """

    @patch.dict(
        os.environ,
        {
            "ZOHO_CLIENT_ID": "id",
            "ZOHO_CLIENT_SECRET": "secret",
            "ZOHO_REFRESH_TOKEN": "refresh",
            "ZOHO_ORGANIZATION_ID": "org1",
            "REPORT_TZ": "UTC",
            "RATE_LIMIT_PER_MINUTE": "6000",
            "REPORT_AS_OF_DATE": "2026-09-13",  # a Sunday (ISO week 37)
        },
        clear=True,
    )
    @patch("sales_by_item_common.requests.Session")
    @patch("sales_by_item_common.requests.post")
    def test_main_writes_expected_file_with_wk_range_filename(self, mock_post, mock_session_cls):
        mock_post.return_value = _mock_response({"access_token": "tok", "api_domain": "https://api"})

        session = MagicMock()
        mock_session_cls.return_value = session

        items_page = _mock_response({
            "items": [{"item_id": "item-1"}],
            "page_context": {"has_more_page": False},
        })
        item_detail = _mock_response({
            "item": {
                "item_id": "item-1",
                "sku": "SKU1",
                "name": "Widget",
                "brand": "Acme",
                "item_type": "inventory",
            }
        })
        invoices_page = _mock_response({
            "invoices": [{"invoice_id": "inv-1", "date": "2026-09-10", "status": "paid"}],
            "page_context": {"has_more_page": False},
        })
        invoice_detail = _mock_response({
            "invoice": {
                "invoice_id": "inv-1",
                "date": "2026-09-10",
                "line_items": [{"item_id": "item-1", "quantity": 7}],
            }
        })

        def get_side_effect(url, params=None, headers=None, timeout=None):
            if url.endswith("/inventory/v1/items"):
                return items_page
            if url.endswith("/inventory/v1/items/item-1"):
                return item_detail
            if url.endswith("/inventory/v1/invoices"):
                return invoices_page
            if url.endswith("/inventory/v1/invoices/inv-1"):
                return invoice_detail
            raise AssertionError(f"Unexpected GET {url}")

        session.get.side_effect = get_side_effect

        # Reference date 2026-09-13 (a Sunday, ISO week 37) with a 36-period
        # window ending on that week means the oldest week is ISO week 2.
        output_path = "reports/Sales_by_item(WK2-WK37).xlsx"
        self.assertFalse(os.path.exists(output_path), "test fixture path should not pre-exist")

        try:
            report.main()
            self.assertTrue(os.path.exists(output_path))

            import openpyxl

            workbook = openpyxl.load_workbook(output_path)
            sheet = workbook.active
            header = [c.value for c in sheet[1]]
            self.assertEqual(header[:3], ["SKU", "Item Name", "Brand"])
            self.assertEqual(len(header), 3 + 36)  # 36-period window
            self.assertEqual(header[3], "WK 2")  # oldest week first
            self.assertEqual(header[-1], "WK 37")  # week ending on the run's own Sunday
            self.assertEqual(sheet.cell(row=2, column=1).value, "SKU1")
            self.assertEqual(sheet.cell(row=2, column=sheet.max_column).value, 7)
        finally:
            if os.path.exists(output_path):
                os.remove(output_path)


if __name__ == "__main__":
    unittest.main()
